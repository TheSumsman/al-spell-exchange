"""Check that what is committed is safe to publish and matches a fresh build.

    python scripts/check_artifacts.py            # the working tree (CI)
    python scripts/check_artifacts.py --staged   # what is about to be committed

The --staged form is what the pre-commit hook runs (scripts/hooks/pre-commit).
It writes the INDEX out to a temporary directory and checks that, so it judges
exactly what the commit would contain -- the staged build/ against the STAGED
scripts. Checking against the working tree instead would wrongly pass, or
wrongly fail, a commit that stages only part of a change.

This repository is public and its owner is pseudonymous, so two kinds of
mistake matter here, and both have happened or nearly happened:

  * Leaks. An Office lock file (~$*.xlsx) holds the real name of whoever had
    the workbook open; one was committed and pushed. A rendered table-tent PDF,
    a table tent built with --url, or a Form script built with
    --spreadsheet-id would publish an event's live links. A workbook that has
    been opened in Excel or round-tripped through Google Sheets carries its
    editor's name in its metadata, and may carry real player data.
  * Staleness. build/ is committed so a clone works without running anything,
    which means a script change without a rebuild ships a workbook that does
    not match its own code.

Every check below compares against a FRESH render from the generators, so the
committed artifacts must be exactly what the scripts produce -- no event data,
no live links, no hand edits.

Failure messages say WHERE, never WHAT: CI logs on a public repo are public, so
printing an offending cell value or URL would leak it a second time.
"""
import argparse
import contextlib
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_workbook as bw  # noqa: E402
import make_form_script  # noqa: E402
import make_polish_script  # noqa: E402
import make_table_tent  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

WORKBOOK = "build/SpellExchange.xlsx"
FORM_GS = "build/CreateSpellExchangeForm.gs"
POLISH_GS = "build/PolishSpellExchangeSheet.gs"
TENT = "build/table-tent.html"

# Tracked paths that must never be committed, and why. Most are gitignored as
# well, but `git add -f` and a careless `git add -A` both get past .gitignore.
FORBIDDEN = [
    (re.compile(r"(^|/)~\$"), "no Office lock files",
     "an Office lock file holds the real name of whoever had the document open"),
    (re.compile(r"\.pdf$", re.I), "no PDFs",
     "a rendered table tent keeps the form link in a PDF link annotation"),
    (re.compile(r"^TheMasterSpellbook/"), "no saved D&D Beyond pages",
     "copyrighted content, never published"),
    (re.compile(r"^(BACKLOG|RECORDING)\.md$"), "no local scratch files",
     "BACKLOG.md and RECORDING.md are private notes"),
]
OFFICE = re.compile(r"\.(xlsx|xlsm|xls|docx|doc|pptx|ppt)$", re.I)


class Checker:
    def __init__(self, snapshot):
        self.snapshot = snapshot      # ROOT is a written-out copy of the index
        self.failures = []

    def read(self, path):
        with open(os.path.join(ROOT, path), "rb") as fh:
            return fh.read()

    def text(self, path):
        # Git may check text files out with CRLF on Windows; the generators
        # write LF. Line endings are not what this is checking.
        return self.read(path).decode("utf-8").replace("\r\n", "\n")

    def check(self, label, ok, fix=""):
        print("  %-58s %s" % (label, "ok" if ok else "FAIL"))
        if not ok:
            self.failures.append((label, fix))


def tracked_files(c):
    if c.snapshot:
        # checkout-index wrote exactly the index's files, so the snapshot's
        # contents ARE the next commit's file list.
        return sorted(os.path.relpath(os.path.join(d, f), ROOT).replace(os.sep, "/")
                      for d, dirs, files in os.walk(ROOT)
                      if "__pycache__" not in d for f in files)
    out = subprocess.run(["git", "ls-files", "-z"], cwd=ROOT,
                         capture_output=True, check=True).stdout
    return [p for p in out.decode("utf-8").split("\0") if p]


def check_tracked(c):
    print("\nNothing that must never be published")
    files = tracked_files(c)
    for pattern, label, why in FORBIDDEN:
        hits = [p for p in files if pattern.search(p)]
        c.check(label, not hits,
                "%s: %s. Remove with `git rm --cached`." % (", ".join(hits), why))
    stray = [p for p in files if OFFICE.search(p) and p != WORKBOOK]
    c.check("no Office files except %s" % WORKBOOK, not stray,
            "%s: binary Office files carry author names and paths in their "
            "metadata. Only the generated workbook belongs in the repo."
            % ", ".join(stray))


def check_links(c):
    print("\nNo event links baked into the committed builds")
    tent = c.text(TENT)
    c.check("table tent has no form link",
            "<a href" not in tent and "data:image/png" not in tent,
            "%s was built with --url. Re-run `python scripts/make_table_tent.py` "
            "with no arguments." % TENT)
    gs = c.text(FORM_GS)
    m = re.search(r'^var SPREADSHEET_ID = (.*);$', gs, re.M)
    c.check("Form script has no spreadsheet id", bool(m) and m.group(1) == '""',
            "%s was built with --spreadsheet-id. Re-run "
            "`python scripts/make_form_script.py` with no arguments." % FORM_GS)


def check_workbook_metadata(c):
    print("\nWorkbook has never been opened anywhere")
    with zipfile.ZipFile(io.BytesIO(c.read(WORKBOOK))) as z:
        core = z.read("docProps/core.xml").decode("utf-8")
        app = z.read("docProps/app.xml").decode("utf-8")
    creator = re.search(r"<dc:creator[^>]*>([^<]*)<", core)
    c.check("creator is openpyxl",
            bool(creator) and creator.group(1) == "openpyxl",
            "%s was saved by something other than build_workbook.py, and its "
            "creator field may name a person. Rebuild it." % WORKBOOK)
    c.check("no lastModifiedBy",
            "lastModifiedBy" not in core,
            "%s has been saved by Excel or Sheets, which records the editor's "
            "name. Rebuild it." % WORKBOOK)
    c.check("saved by openpyxl", "Openpyxl" in app,
            "%s was last written by another application. Rebuild it." % WORKBOOK)


def workbook_fingerprint(wb):
    """Everything about a workbook that the generator decides, cell by cell."""
    fp = {"sheets": [ws.title for ws in wb.worksheets]}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells[cell.coordinate] = cell.value
        fp[ws.title] = {
            "cells": cells,
            "validations": sorted((dv.type or "", str(dv.sqref),
                                   dv.formula1 or "", dv.formula2 or "")
                                  for dv in ws.data_validations.dataValidation),
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "hidden_cols": sorted(k for k, d in ws.column_dimensions.items()
                                  if d.hidden),
            "hidden_rows": sorted(k for k, d in ws.row_dimensions.items()
                                  if d.hidden),
            "freeze": ws.freeze_panes,
            "state": ws.sheet_state,
        }
    return fp


def check_fresh(c):
    print("\nbuild/ matches what the scripts generate now")
    rebuild = "Rebuild (see README, 'Rebuilding') and stage build/."
    # In --staged mode the generators are the index's, which may predate the
    # render() this compares against -- say so rather than crash.
    old = [m.__name__ for m in (make_form_script, make_polish_script,
                                make_table_tent) if not hasattr(m, "render")]
    if old:
        c.check("generators can be compared", False,
                "the staged %s predate this checker. Stage the current "
                "scripts/make_*.py with it."
                % ", ".join(old))
        return
    for path, render in ((FORM_GS, make_form_script.render),
                         (POLISH_GS, make_polish_script.render),
                         (TENT, make_table_tent.render)):
        c.check("%s is current" % path, c.text(path) == render(), rebuild)

    with tempfile.TemporaryDirectory() as tmp:
        fresh_path = os.path.join(tmp, "fresh.xlsx")
        with contextlib.redirect_stdout(io.StringIO()):
            bw.build(out=fresh_path)
        fresh = workbook_fingerprint(load_workbook(fresh_path))
    committed = workbook_fingerprint(load_workbook(io.BytesIO(c.read(WORKBOOK))))

    # Report coordinates only -- a differing cell may hold a player's name.
    where = []
    if committed["sheets"] != fresh["sheets"]:
        where.append("the tab list")
    for title in fresh["sheets"]:
        a, b = committed.get(title), fresh[title]
        if a is None:
            continue
        keys = set(a["cells"]) | set(b["cells"])
        where += ["%s!%s" % (title, k) for k in sorted(keys)
                  if a["cells"].get(k) != b["cells"].get(k)]
        where += ["%s (%s)" % (title, part) for part in
                  ("validations", "merged", "hidden_cols", "hidden_rows",
                   "freeze", "state") if a[part] != b[part]]
    c.check("%s is current and holds no data" % WORKBOOK, not where,
            "differs at %s%s. If those cells hold player data, this workbook "
            "has been used at an event. %s"
            % (", ".join(where[:8]),
               " and %d more" % (len(where) - 8) if len(where) > 8 else "",
               rebuild))


def check_form_options(c):
    # A hand edit to the CSV is the documented way to change the spell list,
    # and these lists are easy to forget -- they are not under build/.
    print("\ndata/form-options/ matches data/wizard-spells.csv")
    if not hasattr(make_form_script, "render_form_options"):
        c.check("form option lists can be compared", False,
                "the staged scripts/make_form_script.py predates this checker. "
                "Stage the current one with it.")
        return
    stale = [path for path, text in make_form_script.render_form_options().items()
             if not os.path.exists(os.path.join(ROOT, path)) or c.text(path) != text]
    c.check("form option lists match the spell list", not stale,
            "%s out of step with the CSV. Re-run "
            "`python scripts/make_form_script.py` and stage data/form-options/."
            % ", ".join(stale))


def check_index(quiet):
    """Write the index out and run this checker inside it.

    This file is copied in over the staged copy, so the checks themselves are
    the current ones even while the checker is being changed; everything it
    checks -- generators, CSV, build/ -- is the index's.
    """
    with tempfile.TemporaryDirectory() as tmp:
        prefix = tmp.replace(os.sep, "/") + "/"
        subprocess.run(["git", "checkout-index", "--all", "--prefix=" + prefix],
                       cwd=ROOT, check=True)
        dest = os.path.join(tmp, "scripts", "check_artifacts.py")
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        shutil.copyfile(os.path.abspath(__file__), dest)
        cmd = [sys.executable, dest, "--snapshot"] + (["--quiet"] if quiet else [])
        return subprocess.run(cmd, cwd=tmp).returncode


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--staged", action="store_true",
                    help="check the index (what is about to be committed)")
    ap.add_argument("--quiet", action="store_true",
                    help="print the report only if something fails")
    ap.add_argument("--snapshot", action="store_true", help=argparse.SUPPRESS)
    args = ap.parse_args()
    if args.staged:
        return check_index(args.quiet)
    c = Checker(snapshot=args.snapshot)

    report = io.StringIO()
    with contextlib.redirect_stdout(report if args.quiet else sys.stdout):
        print("Checking %s" % ("the staged commit" if args.snapshot
                               else "the working tree"))
        check_tracked(c)
        check_links(c)
        check_workbook_metadata(c)
        check_fresh(c)
        check_form_options(c)
    if args.quiet and c.failures:
        print(report.getvalue(), end="")

    if c.failures:
        print("\nFAILED")
        for label, fix in c.failures:
            print("\n- %s\n  %s" % (label, fix))
        return 1
    if not args.quiet:
        print("\nALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
